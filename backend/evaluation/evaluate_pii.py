from __future__ import annotations

import ast
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Expected location:
# C:\FYP\backend\evaluation\evaluate_pii.py
THIS_FILE = Path(__file__).resolve()
EVALUATION_DIR = THIS_FILE.parent
BACKEND_ROOT = EVALUATION_DIR.parent

if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from app.services.anonymization_service import anonymize_text
from app.services.pii_detector import detect_pii

DATASET_PATH = EVALUATION_DIR / "Secura_Evaluation_Dataset.xlsx"
RESULTS_PATH = EVALUATION_DIR / "Secura_Evaluation_Results.xlsx"
SUMMARY_TEXT_PATH = EVALUATION_DIR / "evaluation_summary.txt"

SUPPORTED_TYPES = ("NAME", "EMAIL", "PHONE", "MALAYSIAN_IC")
LABEL_MAP = {
    "name": "NAME",
    "email": "EMAIL",
    "phone": "PHONE",
    "malaysian_ic": "MALAYSIAN_IC",
    "malaysian ic": "MALAYSIAN_IC",
    "ic": "MALAYSIAN_IC",
}


def normalize_label(label: str) -> str:
    raw = str(label).strip()
    mapped = LABEL_MAP.get(raw.lower())
    if mapped:
        return mapped
    return raw.upper().replace(" ", "_")


def safe_div(numerator: int, denominator: int) -> float | None:
    if denominator == 0:
        return None
    return numerator / denominator


def metric_values(tp: int, fp: int, fn: int) -> tuple[float | None, float | None, float | None]:
    precision = safe_div(tp, tp + fp)
    recall = safe_div(tp, tp + fn)

    if precision is None or recall is None or (precision + recall) == 0:
        f1 = None if precision is None or recall is None else 0.0
    else:
        f1 = 2 * precision * recall / (precision + recall)

    return precision, recall, f1


def pct(value: float | None) -> str:
    return "N/A" if value is None else f"{value * 100:.2f}%"


def parse_ground_truth(raw_value: object) -> list[dict]:
    if raw_value is None or str(raw_value).strip() == "":
        return []

    parsed = ast.literal_eval(str(raw_value))
    output = []

    for item in parsed:
        if not isinstance(item, (tuple, list)) or len(item) != 3:
            raise ValueError(f"Invalid True Predictions item: {item!r}")

        start, end, label = item
        output.append({
            "start": int(start),
            "end": int(end),
            "type": normalize_label(str(label)),
        })

    return output


def normalize_prediction(item: dict) -> dict | None:
    try:
        start = int(item["start"])
        end = int(item["end"])
        pii_type = normalize_label(str(item["type"]))
    except (KeyError, TypeError, ValueError):
        return None

    if pii_type not in SUPPORTED_TYPES:
        return None

    return {
        "start": start,
        "end": end,
        "type": pii_type,
        "value": str(item.get("value", "")),
    }


def entity_key(entity: dict) -> tuple[int, int, str]:
    return int(entity["start"]), int(entity["end"]), str(entity["type"])


def format_entities(entities: list[dict], text: str) -> str:
    formatted = []
    for entity in entities:
        start = int(entity["start"])
        end = int(entity["end"])
        pii_type = str(entity["type"])
        value = text[start:end]
        formatted.append(f"({start}, {end}, {pii_type}, {value!r})")
    return "[" + ", ".join(formatted) + "]"


def count_removed_ground_truth(text: str, expected: list[dict], anonymized_text: str) -> tuple[int, int]:
    """Count expected PII occurrences that no longer remain after anonymization."""
    grouped = Counter()

    for entity in expected:
        value = text[entity["start"]:entity["end"]]
        if entity["type"] in {"NAME", "EMAIL"}:
            grouped[(entity["type"], value.casefold())] += 1
        else:
            grouped[(entity["type"], value)] += 1

    removed_total = 0
    expected_total = sum(grouped.values())

    for (pii_type, value), expected_count in grouped.items():
        haystack = anonymized_text.casefold() if pii_type in {"NAME", "EMAIL"} else anonymized_text
        remaining = haystack.count(value)
        removed = max(0, expected_count - min(expected_count, remaining))
        removed_total += removed

    return removed_total, expected_total


def load_dataset() -> list[dict]:
    if not DATASET_PATH.exists():
        raise FileNotFoundError(f"Dataset not found: {DATASET_PATH}")

    workbook = load_workbook(DATASET_PATH, data_only=True, read_only=True)

    try:
        worksheet = workbook["Testing_Set"] if "Testing_Set" in workbook.sheetnames else workbook.worksheets[0]
        rows = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if not rows:
        raise ValueError("Evaluation dataset is empty")

    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    required = {"Text", "True Predictions"}
    missing = required - set(headers)
    if missing:
        raise ValueError(f"Missing required dataset columns: {sorted(missing)}")

    text_index = headers.index("Text")
    truth_index = headers.index("True Predictions")

    dataset = []
    for row_number, values in enumerate(rows[1:], start=2):
        text = values[text_index] if text_index < len(values) else None
        truth = values[truth_index] if truth_index < len(values) else None

        if text is None or str(text).strip() == "":
            continue

        text = str(text)
        expected = parse_ground_truth(truth)

        # Validate ground-truth spans against the actual text.
        for entity in expected:
            if entity["start"] < 0 or entity["end"] > len(text) or entity["start"] >= entity["end"]:
                raise ValueError(
                    f"Invalid ground-truth span on Excel row {row_number}: {entity}"
                )

        dataset.append({
            "row_number": row_number,
            "text": text,
            "expected": expected,
        })

    return dataset


def style_workbook(workbook: Workbook) -> None:
    header_fill = PatternFill("solid", fgColor="6B2D1A")
    header_font = Font(color="FFFFFF", bold=True)
    sub_fill = PatternFill("solid", fgColor="F4E7D3")

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        for col_idx, column_cells in enumerate(worksheet.columns, start=1):
            max_length = 0
            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, min(len(value), 80))
            worksheet.column_dimensions[get_column_letter(col_idx)].width = min(max(max_length + 2, 10), 45)

        if worksheet.title == "Summary":
            for cell in worksheet[2]:
                cell.fill = sub_fill


def write_results(summary_rows: list[dict], detail_rows: list[dict], metadata: dict) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"

    summary_headers = ["PII Type", "TP", "FP", "FN", "Precision", "Recall", "F1-score"]
    summary.append(summary_headers)

    for item in summary_rows:
        summary.append([
            item["pii_type"],
            item["tp"],
            item["fp"],
            item["fn"],
            item["precision"],
            item["recall"],
            item["f1"],
        ])

    summary.append([])
    summary.append(["Evaluation Information", "Value"])
    for key, value in metadata.items():
        summary.append([key, value])

    details = workbook.create_sheet("Details")
    detail_headers = [
        "Dataset Row",
        "Text",
        "Expected Entities",
        "Predicted Entities",
        "TP",
        "FP",
        "FN",
        "Precision",
        "Recall",
        "F1-score",
        "Expected PII Removed",
        "Expected PII Total",
        "Anonymization Coverage",
    ]
    details.append(detail_headers)

    for item in detail_rows:
        details.append([item[h] for h in detail_headers])

    style_workbook(workbook)
    workbook.save(RESULTS_PATH)


def main() -> None:
    dataset = load_dataset()

    aggregate = defaultdict(lambda: {"tp": 0, "fp": 0, "fn": 0})
    overall_tp = overall_fp = overall_fn = 0
    anonymized_total = expected_anonymization_total = 0
    detail_rows = []

    for case in dataset:
        text = case["text"]
        expected = case["expected"]

        raw_predictions = detect_pii(text)
        predicted = []
        for item in raw_predictions:
            normalized = normalize_prediction(item)
            if normalized is not None:
                predicted.append(normalized)

        expected_counter = Counter(entity_key(entity) for entity in expected)
        predicted_counter = Counter(entity_key(entity) for entity in predicted)

        matched = expected_counter & predicted_counter
        false_positive = predicted_counter - expected_counter
        false_negative = expected_counter - predicted_counter

        tp = sum(matched.values())
        fp = sum(false_positive.values())
        fn = sum(false_negative.values())
        precision, recall, f1 = metric_values(tp, fp, fn)

        overall_tp += tp
        overall_fp += fp
        overall_fn += fn

        for pii_type in SUPPORTED_TYPES:
            aggregate[pii_type]["tp"] += sum(
                count for (_, _, kind), count in matched.items() if kind == pii_type
            )
            aggregate[pii_type]["fp"] += sum(
                count for (_, _, kind), count in false_positive.items() if kind == pii_type
            )
            aggregate[pii_type]["fn"] += sum(
                count for (_, _, kind), count in false_negative.items() if kind == pii_type
            )

        anonymized = anonymize_text(text)["anonymized_text"]
        removed, expected_total = count_removed_ground_truth(text, expected, anonymized)
        anonymized_total += removed
        expected_anonymization_total += expected_total

        detail_rows.append({
            "Dataset Row": case["row_number"],
            "Text": text,
            "Expected Entities": format_entities(expected, text),
            "Predicted Entities": format_entities(predicted, text),
            "TP": tp,
            "FP": fp,
            "FN": fn,
            "Precision": pct(precision),
            "Recall": pct(recall),
            "F1-score": pct(f1),
            "Expected PII Removed": removed,
            "Expected PII Total": expected_total,
            "Anonymization Coverage": pct(safe_div(removed, expected_total)),
        })

    summary_rows = []
    for pii_type in SUPPORTED_TYPES:
        tp = aggregate[pii_type]["tp"]
        fp = aggregate[pii_type]["fp"]
        fn = aggregate[pii_type]["fn"]
        precision, recall, f1 = metric_values(tp, fp, fn)
        summary_rows.append({
            "pii_type": pii_type,
            "tp": tp,
            "fp": fp,
            "fn": fn,
            "precision": pct(precision),
            "recall": pct(recall),
            "f1": pct(f1),
        })

    overall_precision, overall_recall, overall_f1 = metric_values(
        overall_tp, overall_fp, overall_fn
    )
    summary_rows.append({
        "pii_type": "OVERALL_MICRO",
        "tp": overall_tp,
        "fp": overall_fp,
        "fn": overall_fn,
        "precision": pct(overall_precision),
        "recall": pct(overall_recall),
        "f1": pct(overall_f1),
    })

    anonymization_coverage = safe_div(anonymized_total, expected_anonymization_total)
    metadata = {
        "Dataset file": DATASET_PATH.name,
        "Test records": len(dataset),
        "Expected PII occurrences": overall_tp + overall_fn,
        "Evaluation matching": "Exact entity span + entity type",
        "Overall Precision": pct(overall_precision),
        "Overall Recall": pct(overall_recall),
        "Overall F1-score": pct(overall_f1),
        "Anonymization Coverage": pct(anonymization_coverage),
        "PII removed / expected": f"{anonymized_total}/{expected_anonymization_total}",
        "Status": "Pilot evaluation for progress presentation",
    }

    write_results(summary_rows, detail_rows, metadata)

    lines = [
        "SECURA PII PILOT EVALUATION",
        "=" * 32,
        f"Dataset: {DATASET_PATH.name}",
        f"Test records: {len(dataset)}",
        f"Expected PII occurrences: {overall_tp + overall_fn}",
        "Matching rule: exact entity span + entity type",
        "",
        "Per-type detection metrics:",
    ]

    for item in summary_rows[:-1]:
        lines.append(
            f"- {item['pii_type']}: "
            f"Precision={item['precision']} | "
            f"Recall={item['recall']} | "
            f"F1={item['f1']} | "
            f"TP={item['tp']} FP={item['fp']} FN={item['fn']}"
        )

    lines.extend([
        "",
        "Overall (micro-averaged):",
        f"- Precision: {pct(overall_precision)}",
        f"- Recall: {pct(overall_recall)}",
        f"- F1-score: {pct(overall_f1)}",
        "",
        "End-to-end anonymization coverage:",
        f"- Expected PII removed: {anonymized_total}/{expected_anonymization_total}",
        f"- Coverage: {pct(anonymization_coverage)}",
        "",
        "Important:",
        "- This is a pilot evaluation for progress presentation.",
        "- For the final FYP, expand the dataset and independently review labels.",
        "- Precision = proportion of detected entities that are correct.",
        "- Recall = proportion of ground-truth PII entities that are found.",
        "- F1-score = harmonic balance between Precision and Recall.",
    ])

    summary_text = "\n".join(lines)
    SUMMARY_TEXT_PATH.write_text(summary_text, encoding="utf-8")

    print(summary_text)
    print()
    print(f"Results workbook: {RESULTS_PATH}")
    print(f"Summary text:     {SUMMARY_TEXT_PATH}")


if __name__ == "__main__":
    main()
